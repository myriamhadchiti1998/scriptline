"""STAGE 7 — Delivery. Passing scripts are written as a formatted Word document
(.docx, for the copy team / Google Drive import) plus a Markdown twin (for Notion,
GitHub, or diffing). Held scripts get a Markdown hold report instead — they are
never written to the delivered folder.

Layout of every delivered script:
  page 1  – production notes (top, shaded, labelled as videographer helper),
            Hook 1 / Hook 2 / Hook 3, then the main body (Pain › Promise › Proof › CTA)
  page 2  – appendix: QA scorecard, voice scores, claim → proof ledger
"""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from ..clientcfg import ClientConfig
from ..config import settings
from ..schemas import ScriptDraft

INK = RGBColor(0x1F, 0x23, 0x2B)
MUTED = RGBColor(0x6B, 0x72, 0x80)
ACCENT = RGBColor(0x1D, 0x4E, 0x89)
PASS = RGBColor(0x1E, 0x7B, 0x4F)
FAIL = RGBColor(0xB4, 0x23, 0x18)
HOOK_LABELS = ["Hook 1 — pattern interrupt", "Hook 2 — alternate angle", "Hook 3 — third option"]


def hook_labels(d: ScriptDraft) -> list[str]:
    plan = {h.hook: h.interrupt_type.replace("-", " ") for h in d.hook_plan}
    return [f"{lbl} · {plan[f'hook_{n}']}" if f"hook_{n}" in plan else lbl
            for n, lbl in enumerate(HOOK_LABELS, 1)]


def slug(text: str, n: int = 48) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")[:n].strip("-") or "script"


def _shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _widths(table, widths_pt: list[int]) -> None:
    table.autofit = False
    for i, w in enumerate(widths_pt):
        table.columns[i].width = Pt(w)
        for cell in table.columns[i].cells:
            cell.width = Pt(w)


def _run(par, text: str, *, size: float = 11, bold: bool = False, italic: bool = False, color: RGBColor = INK):
    r = par.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return r


def _label(doc, text: str, color: RGBColor = MUTED, space_before: int = 14):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(2)
    _run(p, text.upper(), size=8.5, bold=True, color=color)
    return p


def _para(doc, text: str, size: float = 12, after: int = 8):
    for chunk in [c for c in text.split("\n") if c.strip()]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = 1.25
        _run(p, chunk.strip(), size=size)


def _meta_rows(client: ClientConfig, idea: dict, review: dict) -> list[tuple[str, str]]:
    m = review["metrics"]
    return [
        ("Client", client.name),
        ("Concept", idea["concept_hook"]),
        ("Avatar · pain", f"{client.avatar_label(idea['target_avatar'])} · {idea['pain_addressed']}"),
        ("Angle", idea["angle"]),
        ("Creative · lead · funnel", f"{idea['creative_type']} · {idea['lead_type']} · {idea['funnel_stage']} traffic"),
        ("ICE", f"{idea['ice_total']}  (I {idea['ice_impact']} · C {idea['ice_confidence']} · E {idea['ice_ease']})"),
        ("Est. A-roll length", f"{m['seconds_range'][0]}–{m['seconds_range'][1]} seconds ({m['total_words_range'][0]}–{m['total_words_range'][1]} words)"),
        ("Quality gate", f"PASSED 10/10 · reviewer score {review['score']}/5"),
    ]


def build_docx(path: Path, client: ClientConfig, idea: dict, script_id: int, d: ScriptDraft, review: dict) -> None:
    doc = Document()
    sec = doc.sections[0]
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Pt(64))
    sec.top_margin = sec.bottom_margin = Pt(56)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    p = doc.add_paragraph()
    _run(p, f"{client.name.upper()}  ·  VIDEO AD SCRIPT #{script_id:03d}", size=8.5, bold=True, color=MUTED)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    _run(p, d.title, size=20, bold=True, color=ACCENT)

    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for k, v in _meta_rows(client, idea, review):
        row = t.add_row().cells
        _run(row[0].paragraphs[0], k, size=9, bold=True, color=MUTED)
        _run(row[1].paragraphs[0], v, size=9.5)
    _widths(t, [120, 350])

    # Production notes — top, visually separated, clearly a helper
    if d.production_notes:
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        box = doc.add_table(rows=1, cols=1).cell(0, 0)
        _shade(box, "EEF2F7")
        _run(box.paragraphs[0], "PRODUCTION NOTES — FOR THE VIDEOGRAPHER (NOT SPOKEN)", size=8.5, bold=True, color=ACCENT)
        for n in d.production_notes:
            bp = box.add_paragraph()
            _run(bp, f"•  {n}", size=9.5, color=INK)

    _label(doc, "Hooks — film all three, test in the first 3 seconds", ACCENT, space_before=18)
    for label, hook in zip(hook_labels(d), d.hooks):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        _run(p, label, size=9, bold=True, color=MUTED)
        _para(doc, hook, size=12.5, after=4)

    _label(doc, "Main body", ACCENT, space_before=18)
    for name, text in (("Pain", d.body.pain), ("Promise", d.body.promise), ("Proof", d.body.proof), ("CTA", d.body.cta)):
        _label(doc, name, space_before=8)
        _para(doc, text)

    # Appendix — QA scorecard & claim ledger (kept off the script page)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    p = doc.add_paragraph()
    _run(p, "Appendix — quality gate report", size=15, bold=True, color=ACCENT)
    p = doc.add_paragraph()
    _run(p, f"Passed all 10 checklist items · reviewer {review.get('judge_model', '')} · "
            f"rule: {review['pass_rule']}", size=9, italic=True, color=MUTED)
    qa = doc.add_table(rows=1, cols=4)
    qa.style = "Light Grid Accent 1"
    for c, h in zip(qa.rows[0].cells, ("#", "Checklist item", "Result", "Evidence")):
        _run(c.paragraphs[0], h, size=9, bold=True)
    for it in review["items"]:
        row = qa.add_row().cells
        _run(row[0].paragraphs[0], str(it["id"]), size=9)
        _run(row[1].paragraphs[0], it["name"], size=9)
        _run(row[2].paragraphs[0], f"{'PASS' if it['passed'] else 'FAIL'} · {it['ai_score']}/5", size=9, bold=True,
             color=PASS if it["passed"] else FAIL)
        _run(row[3].paragraphs[0], it["evidence"][:220], size=8.5, color=MUTED)
    _widths(qa, [22, 170, 62, 216])

    v = review["voice"]
    _label(doc, "Voice & rhythm", space_before=14)
    _para(doc, f"Plainspoken {v['plainspoken']}/5 · Unhurried rhythm {v['unhurried_rhythm']}/5 · "
               f"Operational specificity {v['operational_specificity']}/5 · Anti-hype {v['anti_hype']}/5 · "
               f"Avatar language {v['avatar_language']}/5. {v.get('notes', '')}", size=9.5)

    _label(doc, "Claim ledger — every factual claim and its proof-bank source", space_before=10)
    led = doc.add_table(rows=1, cols=2)
    led.style = "Light Grid Accent 1"
    _run(led.rows[0].cells[0].paragraphs[0], "Claim", size=9, bold=True)
    _run(led.rows[0].cells[1].paragraphs[0], "Proof ID", size=9, bold=True)
    for c in d.claims:
        row = led.add_row().cells
        _run(row[0].paragraphs[0], c.claim, size=9)
        _run(row[1].paragraphs[0], c.proof_id, size=9, bold=True)
    _widths(led, [410, 60])

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def build_markdown(client: ClientConfig, idea: dict, script_id: int, d: ScriptDraft, review: dict, *, held: bool = False) -> str:
    L: list[str] = []
    status = "HELD — NOT DELIVERED" if held else "Delivered"
    L += [f"# {d.title}", "", f"**{client.name}** · Script #{script_id:03d} · {status} · "
          f"{datetime.now().strftime('%Y-%m-%d %H:%M')}", ""]
    L += ["| | |", "|---|---|"]
    for k, v in _meta_rows(client, idea, review)[:-1]:
        L.append(f"| **{k}** | {v} |")
    gate = "PASSED 10/10" if review["passed"] else f"FAILED items {review['failed_items']}"
    L += [f"| **Quality gate** | {gate} · reviewer score {review['score']}/5 |", ""]
    if d.production_notes:
        L += ["> **Production notes — for the videographer (not spoken)**", ">"]
        L += [f"> - {n}" for n in d.production_notes]
        L.append("")
    L += ["## Hooks", ""]
    for label, hook in zip(hook_labels(d), d.hooks):
        L += [f"**{label}**", "", hook, ""]
    L += ["## Main body", ""]
    for name, text in (("Pain", d.body.pain), ("Promise", d.body.promise), ("Proof", d.body.proof), ("CTA", d.body.cta)):
        L += [f"**{name.upper()}**", "", text, ""]
    L += ["---", "", "## Appendix — quality gate report", "", "| # | Item | Result | Evidence / reason |", "|---|---|---|---|"]
    for it in review["items"]:
        res = f"{'✅ PASS' if it['passed'] else '❌ FAIL'} {it['ai_score']}/5"
        why = it["evidence"] if it["passed"] else "; ".join(it["fail_reasons"])
        L.append(f"| {it['id']} | {it['name']} | {res} | {why.replace('|', '/')[:300]} |")
    v = review["voice"]
    L += ["", f"**Voice:** plainspoken {v['plainspoken']}/5 · rhythm {v['unhurried_rhythm']}/5 · specificity "
              f"{v['operational_specificity']}/5 · anti-hype {v['anti_hype']}/5 · avatar language {v['avatar_language']}/5",
          "", "**Claim ledger**", "", "| Claim | Proof ID |", "|---|---|"]
    L += [f"| {c.claim.replace('|', '/')} | {c.proof_id} |" for c in d.claims]
    if review.get("top_fix"):
        L += ["", f"**Reviewer's top fix:** {review['top_fix']}"]
    return "\n".join(L) + "\n"


def deliver(client: ClientConfig, idea: dict, script_id: int, d: ScriptDraft, review: dict) -> tuple[str, str]:
    base = settings.output_dir / client.id / "delivered" / f"{script_id:03d}_{slug(d.title)}"
    docx_path, md_path = base.with_suffix(".docx"), base.with_suffix(".md")
    build_docx(docx_path, client, idea, script_id, d, review)
    md_path.write_text(build_markdown(client, idea, script_id, d, review), encoding="utf-8")
    return str(docx_path), str(md_path)


def write_hold_report(client: ClientConfig, idea: dict, script_id: int, d: ScriptDraft, review: dict) -> str:
    path = settings.output_dir / client.id / "held" / f"{script_id:03d}_{slug(d.title)}_HELD.md"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(build_markdown(client, idea, script_id, d, review, held=True), encoding="utf-8")
    return str(path)
